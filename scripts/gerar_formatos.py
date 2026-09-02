#!/usr/bin/env python3
"""Gera os modelos binarios (.docx e .xlsx) a partir das fontes em produto/principal/formatos/.

A fonte e o markdown/yaml; o binario e artefato regeneravel.
Uso: python scripts/gerar_formatos.py
"""
import re
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
FORMATOS = RAIZ / "produto" / "principal" / "formatos"


def falha(msg: str) -> None:
    print(f"ERRO: {msg}")
    sys.exit(1)


def ler_tabelas_markdown(texto: str):
    """Extrai todas as tabelas markdown como listas de linhas (listas de celulas)."""
    tabelas, atual = [], []
    for linha in texto.splitlines():
        if linha.strip().startswith("|"):
            celulas = [c.strip() for c in linha.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in celulas):
                continue  # linha separadora
            atual.append(celulas)
        elif atual:
            tabelas.append(atual)
            atual = []
    if atual:
        tabelas.append(atual)
    return tabelas


def gerar_docx():
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt
    except ImportError:
        falha("python-docx nao instalado (pip install python-docx)")

    fonte = FORMATOS / "esquema_planejamento_tabela.md"
    if not fonte.exists():
        falha(f"fonte nao encontrada: {fonte}")
    texto = fonte.read_text(encoding="utf-8")
    tabelas = ler_tabelas_markdown(texto)
    if len(tabelas) < 2:
        falha("esquema_planejamento_tabela.md deveria ter 2 tabelas (cabecalho e principal)")

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width

    doc.add_heading("Planejamento bimestral", level=1)
    doc.add_paragraph("Horatividade — método A.U.L.A.")

    # tabela de identificacao (pula linha de cabecalho "Campo | Valor")
    ident = tabelas[0][1:]
    t1 = doc.add_table(rows=len(ident), cols=2)
    t1.style = "Table Grid"
    for i, (campo, valor) in enumerate(ident):
        t1.cell(i, 0).text = campo
        t1.cell(i, 1).text = valor
        for run in t1.cell(i, 0).paragraphs[0].runs:
            run.bold = True

    doc.add_paragraph("")

    # tabela principal
    principal = tabelas[1]
    t2 = doc.add_table(rows=len(principal), cols=len(principal[0]))
    t2.style = "Table Grid"
    for j, cel in enumerate(principal[0]):
        c = t2.cell(0, j)
        c.text = cel
        for run in c.paragraphs[0].runs:
            run.bold = True
    for i, linha in enumerate(principal[1:], start=1):
        for j, cel in enumerate(linha):
            t2.cell(i, j).text = cel

    for par in doc.paragraphs:
        for run in par.runs:
            run.font.size = Pt(10)

    doc.add_paragraph("Horatividade — método A.U.L.A. · modelo v0.1.0")

    destino = FORMATOS / "modelo_planejamento.docx"
    doc.save(destino)
    print(f"gerado: {destino.relative_to(RAIZ)}")


def parse_yaml_simples(caminho: Path):
    """Parser minimo para o esquema_planilha_por_aula.yaml (subconjunto conhecido)."""
    colunas, atual = [], None
    em_colunas = False
    em_validacao = False
    for bruto in caminho.read_text(encoding="utf-8").splitlines():
        linha = bruto.split("#", 1)[0].rstrip() if not bruto.strip().startswith("#") else ""
        if not linha.strip():
            continue
        if linha.startswith("colunas:"):
            em_colunas = True
            continue
        if em_colunas and not linha.startswith(" ") and not linha.startswith("-"):
            em_colunas = False  # saiu do bloco colunas (ex.: nota:)
        if not em_colunas:
            continue
        stripped = linha.strip()
        if stripped.startswith("- nome:"):
            atual = {"nome": stripped.split(":", 1)[1].strip().strip('"'), "validacao": []}
            colunas.append(atual)
            em_validacao = False
        elif stripped.startswith("largura:") and atual is not None:
            atual["largura"] = int(stripped.split(":", 1)[1].strip())
            em_validacao = False
        elif stripped.startswith("formato:") and atual is not None:
            atual["formato"] = stripped.split(":", 1)[1].strip().strip('"')
            em_validacao = False
        elif stripped.startswith("validacao:") and atual is not None:
            resto = stripped.split(":", 1)[1].strip()
            em_validacao = True
            if resto.startswith("["):
                atual["validacao"] = [v.strip().strip('"') for v in resto.strip("[]").split(",")]
                em_validacao = False
        elif stripped.startswith('- "') and em_validacao and atual is not None:
            atual["validacao"].append(stripped.lstrip("- ").strip('"'))
    return colunas


def gerar_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        falha("openpyxl nao instalado (pip install openpyxl)")

    fonte = FORMATOS / "esquema_planilha_por_aula.yaml"
    if not fonte.exists():
        falha(f"fonte nao encontrada: {fonte}")
    colunas = parse_yaml_simples(fonte)
    if not colunas:
        falha("nenhuma coluna lida do yaml")

    wb = Workbook()
    ws = wb.active
    ws.title = "Planejamento"
    negrito = Font(bold=True)
    fundo = PatternFill("solid", fgColor="F2F2F2")

    for j, col in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=j, value=col["nome"])
        cel.font = negrito
        cel.fill = fundo
        ws.column_dimensions[get_column_letter(j)].width = col.get("largura", 15)

    ws.freeze_panes = "A2"
    N_LINHAS = 80

    for j, col in enumerate(colunas, start=1):
        letra = get_column_letter(j)
        if col.get("formato") == "dd/mm":
            for i in range(2, N_LINHAS + 2):
                ws.cell(row=i, column=j).number_format = "DD/MM"
        if col["validacao"]:
            formula = '"' + ",".join(col["validacao"]) + '"'
            if len(formula) <= 255:  # limite do Excel para lista inline
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                dv.error = "Use um dos valores aceitos pelo template."
                dv.errorTitle = "Valor fora da lista"
                ws.add_data_validation(dv)
                dv.add(f"{letra}2:{letra}{N_LINHAS + 1}")
            else:
                print(f"aviso: validacao da coluna '{col['nome']}' excede 255 caracteres; pulada")

    destino = FORMATOS / "modelo_planilha.xlsx"
    wb.save(destino)
    print(f"gerado: {destino.relative_to(RAIZ)}")


def main():
    gerar_docx()
    gerar_xlsx()
    print("ok: formatos gerados")


if __name__ == "__main__":
    main()
